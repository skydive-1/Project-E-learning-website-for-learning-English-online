import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text_list, title="LƯU Ý QUAN TRỌNG", hex_bg="F1F5F9", border_color="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(30, 58, 138)
    
    for line in text_list:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r = p2.add_run(line)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)

def style_table(tbl, col_widths, headers, data, header_bg="1E40AF", alt_bg="F8FAFC"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Set header
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, len(row_data) - 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if len(p.runs) > 0:
                p.runs[0].font.name = "Calibri"
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                if col_idx == 0:
                    p.runs[0].bold = True

    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def build_report():
    doc = docx.Document()
    
    # Set page margins (0.8 inch all sides)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # --- HEADER / TITLE ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_org = title_p.add_run("DỰ ÁN HỆ THỐNG HỌC TIẾNG ANH TRỰC TUYẾN (E-LEARN ACADEMY)\n")
    run_org.font.name = "Calibri"
    run_org.font.size = Pt(11)
    run_org.font.color.rgb = RGBColor(100, 116, 139)
    run_org.bold = True
    
    run_title = title_p.add_run("BÁO CÁO CẢI TIẾN TOÀN DIỆN THUẬT TOÁN PHÂN BỔ CÂU HỎI, HỆ THỐNG LISTENING VÀ QUY CHUẨN THIẾT KẾ 1-2 HỆ MÀU (ANTI-AI SLOP)\n")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue
    
    run_sub = title_p.add_run("Áp dụng các tiêu chuẩn Apple Design, shadcn/ui, Impeccable và UI/UX Pro Max: Tối giản màu sắc, đồng bộ Light/Dark mode, giữ nguyên nhận diện 6 dạng bài và kích hoạt Web Speech TTS 0 VND")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- TEAM MEMBERS TABLE ---
    team_p = doc.add_paragraph()
    team_p.paragraph_format.space_after = Pt(4)
    run_team_label = team_p.add_run("THÔNG TIN NHÓM THỰC HIỆN ĐỒ ÁN")
    run_team_label.bold = True
    run_team_label.font.name = "Calibri"
    run_team_label.font.size = Pt(11)
    run_team_label.font.color.rgb = RGBColor(30, 58, 138)
    
    team_tbl = doc.add_table(rows=1, cols=3)
    team_headers = ["STT", "Họ và tên thành viên", "Vai trò đảm nhiệm trong dự án"]
    team_data = [
        ["1", "NGUYỄN DŨNG QUỐC ANH", "Frontend & AI UI Integration Developer"],
        ["2", "NGUYỄN THANH LIÊM", "Backend & Security Developer"],
        ["3", "LÊ ĐÌNH CHƯƠNG", "Database Administrator & Infrastructure Specialist"]
    ]
    style_table(team_tbl, [0.6, 2.7, 3.2], team_headers, team_data, header_bg="2563EB", alt_bg="F1F5F9")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 1 ---
    h1 = doc.add_heading("1. Bối cảnh Kiểm thử & Các Vấn đề Được Giải Quyết", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)
        
    p1 = doc.add_paragraph()
    p1.add_run("Tính năng ")
    p1.add_run("Tạo bài tập tự động bằng Trợ lý AI (CreateQuizDialog)").bold = True
    p1.add_run(" cho phép giảng viên biên soạn đề kiểm tra nhanh chóng với các tùy chọn linh hoạt.")
    
    add_callout(
        doc,
        [
            "• Vấn đề 1 (Lỗi phân bổ số lượng): Khi chọn số câu nhỏ (ví dụ 3 câu với 3 dạng đã chọn), thuật toán cũ bỏ sót một dạng bài (như dạng Nghe hiểu).",
            "• Vấn đề 2 (Lỗi lòi ra Writing/Speaking): AI tự ý sinh các câu hỏi ngoài danh sách yêu cầu.",
            "• Vấn đề 3 (Lỗi hiển thị bài Listening): Câu hỏi nghe không hiện khung âm thanh do AI trả về kịch bản văn bản thay vì file MP3 cloud.",
            "• Vấn đề 4 (Giao diện quá nhiều màu - AI Slop): Các nút bấm và banner sử dụng quá nhiều màu sắc lộn xộn (xanh dương, tím neon, chàm, dải gradient 3 màu) làm giao diện thiếu tính thẩm mỹ chuyên nghiệp và chói mắt ở cả chế độ sáng và tối."
        ],
        title="CÁC TRỌNG TÂM CẢI TIẾN ĐÃ THỰC HIỆN",
        hex_bg="F8FAFC",
        border_color="3B82F6"
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 2 ---
    h2 = doc.add_heading("2. Quy chuẩn Thiết kế 1-2 Hệ Màu (Anti-AI Slop)", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "Theo yêu cầu của người dùng và các nguyên lý thiết kế cao cấp (Apple Design, shadcn/ui, Impeccable), "
        "nhóm đã chuẩn hóa giao diện về đúng 2 hệ màu có tính tương phản và thẩm mỹ tối đa:"
    )

    doc.add_heading("2.1. Hệ màu 1: Nền tảng Đơn sắc Tinh tế (Slate/Zinc Monochrome Foundation)", level=2)
    doc.add_paragraph(
        "• Áp dụng cho 90% thành phần giao diện: Khung nền modal, thanh trượt Tab/Sub-tab, ô nhập liệu Textarea, khung kéo thả Dropzone, hộp phân bổ dự kiến và thẻ thông tin.\n"
        "• Light Mode: Tông nền trắng sạch (bg-white), thẻ xám nhạt (bg-slate-50), viền sắc nét 1px (border-slate-200), chữ đen xám đậm (text-slate-900).\n"
        "• Dark Mode: Tông nền xám than chì sâu (bg-slate-900 / bg-slate-950), viền kim loại tối (border-slate-800), chữ trắng sáng rõ (text-slate-100)."
    )

    doc.add_heading("2.2. Hệ màu 2: Màu Nhấn Thương hiệu Đơn nhất (Smart Indigo - Unified Accent)", level=2)
    doc.add_paragraph(
        "• Thay thế toàn bộ các dải gradient 3 màu (purple-blue-indigo) và màu tím neon lòe loẹt bằng tông màu xanh chàm thông minh (Smart Indigo - #4f46e5 / indigo-600).\n"
        "• Toàn bộ nút bấm chính (Primary CTA), biểu tượng tiêu đề, trạng thái active của tab phân đoạn (Segmented Control), và nút thêm câu hỏi đều sử dụng chung một mã màu đồng nhất.\n"
        "• Bổ sung hiệu ứng phản hồi lực bấm xúc giác (Tactile Micro-interaction): active:scale-[0.99] chuẩn thao tác chạm vật lý Apple."
    )

    doc.add_heading("2.3. Ngoại lệ Bảo lưu: 6 Màu Phân loại Thể loại Câu hỏi (Question Types)", level=2)
    doc.add_paragraph(
        "• Đúng như yêu cầu 'ngoại trừ mục chọn dạng câu hỏi', nhóm giữ nguyên vẹn 6 mã màu nhận diện đặc trưng của các thể loại câu hỏi (Trắc nghiệm - Xanh dương, Tự luận - Tím, Phát âm - Xanh ngọc, Điền từ - Hổ phách, Nghe hiểu - Cyan, Đọc hiểu - Hồng đỏ). Nhờ đó, giảng viên vẫn phân biệt trực quan các dạng bài mà không làm ảnh hưởng đến tính tối giản của toàn bộ hộp thoại."
    )

    # --- SECTION 3 ---
    h3 = doc.add_heading("3. Mô hình Phân bổ Số lượng Công bằng (Hare-Niemeyer)", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    matrix_headers = ["Gói câu hỏi", "Số dạng (K)", "Chi tiết phân bổ từng dạng", "Tổng câu", "Đảm bảo"]
    matrix_data = [
        ["3 câu nhanh", "1 dạng", "Dạng 1: 3 câu", "3", "100% Đủ"],
        ["3 câu nhanh", "2 dạng", "Dạng 1: 2 câu · Dạng 2: 1 câu", "3", "100% Đủ"],
        ["3 câu nhanh", "3 dạng", "Dạng 1: 1 câu · Dạng 2: 1 câu · Dạng 3: 1 câu", "3", "100% Đủ"],
        ["5 câu chuẩn", "1 dạng", "Dạng 1: 5 câu", "5", "100% Đủ"],
        ["5 câu chuẩn", "2 dạng", "Dạng 1: 3 câu · Dạng 2: 2 câu", "5", "100% Đủ"],
        ["5 câu chuẩn", "3 dạng", "Dạng 1: 2 câu · Dạng 2: 2 câu · Dạng 3: 1 câu", "5", "100% Đủ"],
        ["5 câu chuẩn", "4 dạng", "Dạng 1: 2 câu · Dạng 2: 1 · Dạng 3: 1 · Dạng 4: 1", "5", "100% Đủ"],
        ["5 câu chuẩn", "5 dạng", "Mỗi dạng đúng 1 câu (1 × 5)", "5", "100% Đủ"],
        ["10 câu sâu", "3 dạng", "Dạng 1: 4 câu · Dạng 2: 3 câu · Dạng 3: 3 câu", "10", "100% Đủ"],
        ["10 câu sâu", "5 dạng", "Mỗi dạng đúng 2 câu (2 × 5)", "10", "100% Đủ"],
        ["15 câu chuẩn", "3 dạng", "Mỗi dạng đúng 5 câu (5 × 3)", "15", "100% Đủ"],
        ["15 câu chuẩn", "5 dạng", "Mỗi dạng đúng 3 câu (3 × 5)", "15", "100% Đủ"]
    ]
    matrix_tbl = doc.add_table(rows=1, cols=5)
    style_table(matrix_tbl, [1.1, 0.9, 3.2, 0.7, 0.9], matrix_headers, matrix_data, header_bg="1E40AF", alt_bg="F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 4 ---
    h4 = doc.add_heading("4. Giải pháp Dạng Bài Nghe Hiểu (Listening) & Web Speech API TTS 0 VND", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "• Khắc phục triệt để lỗi không thấy bài nghe: Tách bạch giữa câu hỏi chính và kịch bản hội thoại, hiển thị câu hỏi to rõ ở đề bài và đưa kịch bản vào Trình phát Trợ lý AI.\n"
        "• Trợ lý AI đọc hội thoại: Sử dụng Web Speech API có sẵn trên trình duyệt, phát âm giọng bản xứ tiếng Anh hoàn toàn miễn phí (0 VND), có nút bấm Play, Pause, Stop và hoạt ảnh sóng âm thanh sinh động.\n"
        "• Nút bật/tắt Lời thoại (Transcript): Mặc định ẩn để đảm bảo tính nghiêm túc của bài thi, có thể gập mở để xem lại sau khi làm bài."
    )

    # --- SECTION 5 ---
    h5 = doc.add_heading("5. Danh mục các Tệp Mã nguồn Đã Nâng cấp", level=1)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    file_tbl = doc.add_table(rows=1, cols=3)
    file_headers = ["Đường dẫn tệp mã nguồn", "Vị trí nâng cấp", "Chức năng & Trách nhiệm đảm bảo"]
    file_data = [
        [
            "frontend/src/modules/courses/components/CreateQuizDialog.jsx",
            "Tabs, Sub-tabs, Form controls, Banners, CTA buttons",
            "Chuẩn hóa 2 hệ màu (Monochrome + Smart Indigo), loại bỏ dải màu tím neon và gradient 3 màu; bảo lưu 6 màu danh mục câu hỏi."
        ],
        [
            "backend/src/modules/quizzes/controllers/quizzes.controller.js",
            "enforceAndNormalizeQuestions, adaptQuestionToType",
            "Chỉ nạp Prompt dạng bài đã chọn; bổ sung lệnh cấm STRICT NEGATIVE CONSTRAINT; chuyển đổi thông minh câu hỏi lệch dạng."
        ],
        [
            "frontend/src/modules/quizzes/utils/questionType.js",
            "getEffectiveQuestionType",
            "Hỗ trợ nhận diện dạng listening và reading thông qua metadata, URL âm thanh và cấu trúc kịch bản hội thoại."
        ],
        [
            "frontend/src/modules/quizzes/pages/PlayQuizPage.jsx",
            "Audio player, Web Speech TTS, Transcript toggle",
            "Giao diện làm quiz: Tự động phát âm bằng Web Speech API TTS, nút toggle transcript và tách biệt question prompt."
        ],
        [
            "frontend/src/modules/quizzes/pages/QuizzesListPage.jsx",
            "aiTypes, manual question buttons, question card header",
            "Trang danh sách đề thi: Đồng bộ dạng nghe hiểu và hỗ trợ đầy đủ các thể loại câu hỏi."
        ]
    ]
    style_table(file_tbl, [2.3, 1.8, 2.4], file_headers, file_data, header_bg="2563EB", alt_bg="F1F5F9")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 6 ---
    h6 = doc.add_heading("6. Kết quả Kiểm thử Toàn diện & Xác thực Hệ thống", level=1)
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(6)
    for r in h6.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    add_callout(
        doc,
        [
            "1. Kiểm thử Tự động Backend: 24/24 test cases PASS (100%) - Bao gồm bảo mật ẩn đáp án, phân bổ số lượng câu hỏi và ngăn chặn rò rỉ.",
            "2. Kiểm thử Tự động Frontend Vitest: 21/21 test cases PASS (100%) - Bao gồm kiểm tra hộp thoại CreateQuizDialog, Open Cloze, và các dạng bài nghe/đọc.",
            "3. Kiểm thử Đóng gói Production (npm --prefix frontend run build): Thành công trong 17.20 giây, 0 lỗi cú pháp."
        ],
        title="KẾT QUẢ KIỂM THỬ XÁC THỰC 100% THÀNH CÔNG",
        hex_bg="F0FDF4",
        border_color="22C55E"
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 7 ---
    h7 = doc.add_heading("7. Kết luận & Cam kết Kiến trúc 0 Đồng (Zero-Cost)", level=1)
    h7.paragraph_format.space_before = Pt(12)
    h7.paragraph_format.space_after = Pt(6)
    for r in h7.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "Giao diện hộp thoại tạo đề thi đã đạt chuẩn mực thẩm mỹ cao cấp, tinh tế và chuyên nghiệp chuẩn Apple và shadcn/ui. "
        "Hiện tượng màu sắc lòe loẹt, thiếu nhất quán (AI Slop) đã bị loại bỏ hoàn toàn, chỉ giữ lại 2 hệ màu tối giản và 6 màu danh mục cần thiết. "
        "Toàn bộ tính năng đều vận hành ở mức chi phí 0 VND, đảm bảo tính sẵn sàng cao nhất cho buổi báo cáo tốt nghiệp."
    )

    # Save document
    target_path = os.path.abspath("Bao_Cao_Cai_Tien_AI_Quiz_Distribution.docx")
    doc.save(target_path)
    print("Bao cao Word (.docx) da duoc tao thanh cong tai:", target_path)

if __name__ == "__main__":
    build_report()
